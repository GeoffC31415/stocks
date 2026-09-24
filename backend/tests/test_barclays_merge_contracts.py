"""Synthetic regression contracts for the reconciled single Barclays pipeline."""

import asyncio
import datetime as dt
import io
from argparse import Namespace
from unittest.mock import AsyncMock

import openpyxl
import pytest
from sqlalchemy import event, func, select
from test_barclays_sync_service import ACCOUNT, DAY, change_cell, pair, row_counts
from test_barclays_sync_service import db as db

from app.models import Instrument, Order
from app.services import barclays_sync_service as service


@pytest.mark.parametrize(
    "section,column,value",
    [
        ("holdings", 1, None),
        ("holdings", 2, None),
        ("holdings", 3, "n/a"),
        ("holdings", 10, None),
        ("holdings", 10, "nan"),
        ("orders", 2, "not a date"),
        ("orders", 5, "Hold"),
        ("orders", 7, None),
        ("orders", 4, "Other ISA"),
    ],
)
async def test_migrated_malformed_rows_leave_no_database_writes(db, section, column, value):
    session, _ = db
    h, o = pair()
    if section == "holdings":
        h = change_cell(h, column, value)
    else:
        o = change_cell(o, column, value)
    with pytest.raises(service.BarclaysPairError):
        await service.import_pair(session, h, o, as_of=DAY)
    assert await row_counts(session) == [0, 0, 0, 0]
    assert await session.scalar(select(func.count()).select_from(Instrument)) == 0


async def test_empty_snapshot_leaves_no_writes(db):
    session, _ = db
    h, o = pair()
    wb = openpyxl.load_workbook(io.BytesIO(h))
    wb.active.delete_rows(4, 2)
    out = io.BytesIO()
    wb.save(out)
    with pytest.raises(service.BarclaysPairError):
        await service.import_pair(session, out.getvalue(), o, as_of=DAY)
    assert await row_counts(session) == [0, 0, 0, 0]


async def test_initial_cash_only_snapshot_requires_operator_review(db):
    session, _ = db
    h, o = pair()
    wb = openpyxl.load_workbook(io.BytesIO(h))
    wb.active.delete_rows(4)
    out = io.BytesIO()
    wb.save(out)
    with pytest.raises(service.BarclaysPairError, match="operator review"):
        await service.import_pair(session, out.getvalue(), o, as_of=DAY)
    assert await row_counts(session) == [0, 0, 0, 0]


async def test_timeout_in_real_pair_import_rolls_back_and_runs_later_steps(
    db, tmp_path, monkeypatch
):
    from app.services import sync_runner as runner

    session, _ = db
    h, o = pair()
    guard = tmp_path / "attempt"
    guard.write_text("uncertain")
    original = service.ingest_parsed_orders

    async def slow_orders(*args, **kwargs):
        result = await original(*args, **kwargs)
        await asyncio.sleep(0.2)
        return result

    monkeypatch.setattr(service, "ingest_parsed_orders", slow_orders)
    monkeypatch.setattr(runner, "FETCH_TIMEOUT_SECONDS", 0.1)

    async def fetch(inbox):
        return service.FetchedPair(h, o, dt.datetime(2026, 9, 23, tzinfo=dt.UTC), guard.unlink)

    later = AsyncMock(return_value=runner.StepResult("Later", "ok"))
    report = await runner.run_sync_all(
        session,
        fetchers=[("Barclays", fetch), ("Later", later)],
        inbox=tmp_path,
        include_trading212=False,
    )
    assert report.steps[0].status == "failed"
    assert "timed out" in report.steps[0].detail
    assert guard.exists()
    later.assert_awaited_once()
    assert await row_counts(session) == [0, 0, 0, 0]


@pytest.mark.parametrize("factory", ["cli", "router"])
@pytest.mark.parametrize("no_fetch,dry_run", [(False, False), (True, False), (False, True)])
async def test_single_barclays_registration(factory, no_fetch, dry_run, db, tmp_path, monkeypatch):
    from app import sync_cli
    from app.fetchers import barclays, hl
    from app.routers import sync
    from app.services.sync_runner import StepResult, run_sync_all

    fetch = AsyncMock(return_value=StepResult("Barclays", "skipped"))
    monkeypatch.setattr(barclays, "fetch", fetch)
    monkeypatch.setattr(hl, "fetch", AsyncMock(return_value=StepResult("HL", "skipped")))
    if factory == "cli":
        registered = sync_cli._fetchers(Namespace(no_fetch=no_fetch, only=None, headed=False))
    else:
        registered = sync._fetchers(not no_fetch)
    assert [name for name, _ in registered].count("Barclays") == (0 if no_fetch else 1)
    await run_sync_all(
        db[0], fetchers=registered, dry_run=dry_run, include_trading212=False, inbox=tmp_path
    )
    assert fetch.await_count == (0 if no_fetch or dry_run else 1)
    assert not hasattr(sync_cli, "_session_steps")
    assert not hasattr(sync, "_session_steps")


@pytest.mark.parametrize("commit", [False, True])
async def test_canonical_services_real_matching_obey_transaction_ownership(db, commit):
    from app.services.import_service import import_holding_snapshot
    from app.services.order_service import ingest_parsed_orders

    session, engine = db
    h, o = pair()
    parsed = service.validate_pair(h, o, as_of=DAY)
    commits = []
    event.listen(engine.sync_engine, "commit", lambda conn: commits.append(True))
    kwargs = {} if commit else {"commit": False}
    await ingest_parsed_orders(
        session, parsed=parsed.orders, file_bytes=o, filename="synthetic.xls", **kwargs
    )
    await import_holding_snapshot(
        session,
        parsed_rows=parsed.holdings,
        as_of_date=DAY,
        file_sha256="a" * 64,
        filename="synthetic.xls",
        **kwargs,
    )
    order = await session.scalar(select(Order))
    assert order.instrument_id is not None
    assert order.match_status == "auto_high"
    assert bool(commits) is commit
    if not commit:
        await session.rollback()
        assert await row_counts(session) == [0, 0, 0, 0]
        assert await session.scalar(select(func.count()).select_from(Instrument)) == 0
    else:
        await session.rollback()
        assert await row_counts(session) == [1, 2, 1, 1]


@pytest.mark.parametrize("cancel", [False, True])
async def test_joined_session_contains_even_legacy_inner_commit(db, monkeypatch, cancel):
    session, engine = db
    commits = []
    event.listen(engine.sync_engine, "commit", lambda conn: commits.append(True))
    original = service.ingest_parsed_orders

    async def fail_after_real_matching(worker, **kwargs):
        await original(worker, **kwargs)
        order = await worker.scalar(select(Order))
        assert order.instrument_id is not None
        await worker.commit()  # A legacy helper must not commit the owning connection.
        if cancel:
            raise asyncio.CancelledError()
        raise RuntimeError("synthetic late failure")

    monkeypatch.setattr(service, "ingest_parsed_orders", fail_after_real_matching)
    h, o = pair()
    with pytest.raises(asyncio.CancelledError if cancel else RuntimeError):
        await service.import_pair(session, h, o, as_of=DAY)
    assert commits == []
    assert await row_counts(session) == [0, 0, 0, 0]
    assert await session.scalar(select(func.count()).select_from(Instrument)) == 0


@pytest.mark.parametrize("day", [1, 23])
async def test_even_sell_orders_cannot_authorize_disappearing_positions(db, day):
    session, _ = db
    h, o = pair()
    # Keep a second security so this exercises disappearance, not empty export rejection.
    wb = openpyxl.load_workbook(io.BytesIO(h))
    wb.active.append(["Other", "OTHER", 1, 1, "GBP", 1, "GBP", 1, 100, 1])
    out = io.BytesIO()
    wb.save(out)
    await service.import_pair(session, out.getvalue(), o, as_of=DAY)
    sells = change_cell(
        change_cell(change_cell(o, 5, "Sell"), 2, dt.datetime(2026, 9, day)), 1, "Other"
    )
    with pytest.raises(service.BarclaysPairError, match="disappeared"):
        await service.import_pair(session, h, sells, as_of=DAY + dt.timedelta(days=1))
    assert await row_counts(session) == [1, 3, 1, 1]


async def test_multiple_account_sheets_rejected_without_writes(db):
    session, _ = db
    h, o = pair()
    wb = openpyxl.load_workbook(io.BytesIO(h))
    wb.copy_worksheet(wb.active).title = "Other account"
    out = io.BytesIO()
    wb.save(out)
    with pytest.raises(service.BarclaysPairError):
        await service.import_pair(session, out.getvalue(), o, as_of=DAY)
    assert await row_counts(session) == [0, 0, 0, 0]


async def test_cancelled_order_with_absent_amount_is_counted_not_imported(db):
    session, _ = db
    h, o = pair()
    o = change_cell(change_cell(o, 3, "Cancelled"), 7, None)
    result = await service.import_pair(session, h, o, as_of=DAY)
    assert result["cancelled_orders"] == 1
    assert result["orders_imported"] == 0
    assert await row_counts(session) == [1, 2, 1, 0]


def test_retired_login_flag_cannot_arm_automation(monkeypatch):
    from app.config import Settings

    monkeypatch.setenv("PORTFOLIO_BARCLAYS_AUTO_LOGIN", "true")
    monkeypatch.delenv("PORTFOLIO_BARCLAYS_AUTOMATION_ENABLED", raising=False)
    config = Settings(_env_file=None)
    assert config.barclays_automation_enabled is False
    assert not hasattr(config, "barclays_auto_login")
    assert config.sync_service_trigger_enabled is False


@pytest.mark.parametrize("failure", [None, "export", "logout", "cancel"])
async def test_sole_browser_path_logs_out_without_ack_or_private_errors(
    tmp_path, monkeypatch, caplog, failure
):
    from contextlib import asynccontextmanager
    from types import SimpleNamespace

    from pydantic import SecretStr

    from app.config import settings
    from app.fetchers import barclays

    monkeypatch.setattr(settings, "browser_profile", tmp_path)
    monkeypatch.setattr(settings, "barclays_automation_enabled", True)
    monkeypatch.setattr(settings, "barclays_expected_account", SecretStr(ACCOUNT))
    for field in ("surname", "membership_number", "passcode", "memorable_word"):
        monkeypatch.setattr(settings, f"barclays_{field}", SecretStr("synthetic"))
    click = AsyncMock(side_effect=RuntimeError("PRIVATE_SENTINEL") if failure == "logout" else None)
    control = SimpleNamespace(count=AsyncMock(return_value=1), first=SimpleNamespace(click=click))
    page = SimpleNamespace(
        url="https://www.investments.barclays.co.uk/en-gb/SubAccount/test/Portfolio",
        get_by_role=lambda *a, **k: control,
    )

    @asynccontextmanager
    async def context(*args, **kwargs):
        yield SimpleNamespace(new_page=AsyncMock(return_value=page))

    monkeypatch.setattr(barclays, "broker_context", context)
    monkeypatch.setattr(barclays, "_login_once", AsyncMock())
    exports = AsyncMock(return_value=pair())
    if failure == "export":
        exports.side_effect = RuntimeError("PRIVATE_SENTINEL")
    elif failure == "cancel":
        exports.side_effect = asyncio.CancelledError()
    monkeypatch.setattr(barclays, "collect_exports", exports)
    if failure == "cancel":
        with pytest.raises(asyncio.CancelledError):
            await barclays.fetch(tmp_path / "inbox")
    else:
        result = await barclays.fetch(tmp_path / "inbox")
        assert isinstance(result, service.FetchedPair) is (failure != "export")
        assert "PRIVATE_SENTINEL" not in str(result)
    click.assert_awaited_once_with(timeout=5000)
    assert barclays.attempt_marker().exists()
    assert "PRIVATE_SENTINEL" not in caplog.text


async def test_logout_never_clicks_an_untrusted_origin():
    from types import SimpleNamespace

    from app.fetchers import barclays

    def forbidden(*args, **kwargs):
        raise AssertionError("Must not inspect controls outside Barclays")

    await barclays._logout(SimpleNamespace(url="https://evil.example", get_by_role=forbidden))
