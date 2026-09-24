"""Atomic Barclays pair contracts, synthetic workbooks and disposable databases."""

import datetime as dt
import io

import openpyxl
import pytest
from sqlalchemy import event, func, select
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

from app.models import AccountAlias, Base, HoldingSnapshot, ImportBatch, Order, OrderImportBatch
from app.services import barclays_sync_service as service

ACCOUNT = "ID0000000-001 (Investment ISA)"
HEADINGS = [
    "Investment",
    "Identifier",
    "Quantity Held",
    "Last Price",
    "Last Price CCY",
    "Value",
    "Value CCY",
    "FX Rate",
    "Last Price (p)",
    "Value (£)",
    "Book Cost",
    "Book Cost CCY",
    "Average FX Rate",
    "Book Cost (£)",
    "% Change",
]
ORDER_HEADINGS = [
    "Investment",
    "Date",
    "Order Status",
    "Account",
    "Buy/Sell",
    "Quantity",
    "Cost/Proceeds",
    "Country",
]
DAY = dt.date(2026, 9, 23)


def workbook(header, rows, account=ACCOUNT):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = account
    ws.append([account])
    ws.append([])
    ws.append(header)
    for row in rows:
        ws.append(row)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def pair():
    holdings = workbook(
        HEADINGS,
        [
            ["Example Plc", "EXM", 2, 5, "GBP", 10, "GBP", 1, 500, 10, 8, "GBP", 1, 8, 25],
            ["Cash", None, None, None, None, 3, "GBP", None, None, 3, None, None, None, None, None],
        ],
    )
    orders = workbook(
        ORDER_HEADINGS,
        [
            [
                "Example Plc",
                dt.datetime(2026, 9, 20),
                "Completed",
                "Investment ISA",
                "Buy",
                2,
                8,
                "UK",
            ]
        ],
    )
    return holdings, orders


def test_validated_pair_preserves_legacy_order_account_identity():
    h, o = pair()
    result = service.validate_pair(h, o, as_of=DAY)
    assert len(result.holdings) == 2
    assert len(result.orders) == 1
    assert result.account_name == ACCOUNT
    assert result.orders[0].account_name == "Investment ISA"
    assert result.cancelled_orders == 0


def change_cell(data, column, value, row=4):
    wb = openpyxl.load_workbook(io.BytesIO(data))
    wb.active.cell(row, column, value=value) if value is not None else setattr(
        wb.active.cell(row, column), "value", None
    )
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


@pytest.mark.parametrize(
    "column,value",
    [(1, None), (2, None), (3, "oops"), (3, "NaN"), (10, "inf"), (10, None), (7, None)],
)
def test_malformed_snapshot_row_rejects_whole_pair(column, value):
    h, o = pair()
    with pytest.raises(service.BarclaysPairError):
        service.validate_pair(change_cell(h, column, value), o, as_of=DAY)


@pytest.mark.parametrize(
    "column,value",
    [
        (1, None),
        (2, "bad date"),
        (3, "unknown"),
        (4, "Other account"),
        (5, "Transfer"),
        (6, "NaN"),
        (7, None),
    ],
)
def test_malformed_order_row_rejects_whole_pair(column, value):
    h, o = pair()
    with pytest.raises(service.BarclaysPairError):
        service.validate_pair(h, change_cell(o, column, value), as_of=DAY)


def test_cancelled_orders_are_explicitly_counted_not_silently_dropped():
    h, o = pair()
    o = change_cell(o, 3, "Cancelled")
    result = service.validate_pair(h, o, as_of=DAY)
    assert result.orders == []
    assert result.cancelled_orders == 1


def test_cash_uses_explicit_gbp_column_when_original_currency_is_blank():
    h, o = pair()
    result = service.validate_pair(change_cell(h, 7, None, row=5), o, as_of=DAY)
    assert result.holdings[1].value_gbp == 3


def test_missing_cash_rejects_incomplete_snapshot():
    h, o = pair()
    wb = openpyxl.load_workbook(io.BytesIO(h))
    wb.active.delete_rows(5)
    out = io.BytesIO()
    wb.save(out)
    with pytest.raises(service.BarclaysPairError):
        service.validate_pair(out.getvalue(), o, as_of=DAY)


@pytest.fixture
async def db():
    engine = create_async_engine("sqlite+aiosqlite:///:memory:")
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    async with async_sessionmaker(engine, expire_on_commit=False)() as session:
        session.add(
            AccountAlias(
                source="barclays_orders",
                source_account_name="Investment ISA",
                canonical_account_name=ACCOUNT,
            )
        )
        await session.commit()
        yield session, engine
    await engine.dispose()


async def row_counts(session):
    return [
        await session.scalar(select(func.count()).select_from(model))
        for model in (ImportBatch, HoldingSnapshot, OrderImportBatch, Order)
    ]


async def test_pair_import_is_idempotent_with_one_physical_commit(db):
    session, engine = db
    commits = []
    event.listen(engine.sync_engine, "commit", lambda conn: commits.append(True))
    h, o = pair()
    result = await service.import_pair(session, h, o, as_of=DAY)
    assert result == {
        "snapshot": "imported",
        "orders": "imported",
        "orders_imported": 1,
        "cancelled_orders": 0,
    }
    assert len(commits) == 1
    assert await row_counts(session) == [1, 2, 1, 1]
    again = await service.import_pair(session, h, o, as_of=DAY)
    assert again["snapshot"] == "unchanged" and again["orders"] == "unchanged"
    assert await row_counts(session) == [1, 2, 1, 1]


async def test_whole_missing_position_requires_review_instead_of_closing(db):
    session, _ = db
    h, o = pair()
    await service.import_pair(session, h, o, as_of=DAY)
    wb = openpyxl.load_workbook(io.BytesIO(h))
    wb.active.delete_rows(4)
    out = io.BytesIO()
    wb.save(out)
    with pytest.raises(service.BarclaysPairError, match="disappear"):
        await service.import_pair(session, out.getvalue(), o, as_of=DAY + dt.timedelta(days=1))
    assert await row_counts(session) == [1, 2, 1, 1]


async def test_missing_country_cannot_duplicate_legacy_order(db):
    session, _ = db
    h, o = pair()
    await service.import_pair(session, h, o, as_of=DAY)
    with pytest.raises(service.BarclaysPairError):
        await service.import_pair(session, h, change_cell(o, 8, None), as_of=DAY)
    assert await row_counts(session) == [1, 2, 1, 1]


async def test_same_day_a_b_a_tracks_latest_observation(db):
    session, _ = db
    h, o = pair()
    await service.import_pair(session, h, o, as_of=DAY)
    await service.import_pair(session, change_cell(h, 10, 20), o, as_of=DAY)
    result = await service.import_pair(session, h, o, as_of=DAY)
    assert result["snapshot"] == "imported"
    assert await row_counts(session) == [3, 6, 1, 1]
    result = await service.import_pair(session, h, o, as_of=DAY)
    assert result["snapshot"] == "unchanged"
    assert await row_counts(session) == [3, 6, 1, 1]


async def test_full_order_label_still_requires_explicit_alias(db):
    session, _ = db
    from sqlalchemy import delete

    await session.execute(delete(AccountAlias))
    await session.commit()
    h, o = pair()
    with pytest.raises(service.BarclaysPairError, match="alias"):
        await service.import_pair(session, h, change_cell(o, 4, ACCOUNT), as_of=DAY)
    assert await row_counts(session) == [0, 0, 0, 0]


async def test_late_order_failure_rolls_back_legacy_matcher_commit(db, monkeypatch):
    session, _ = db
    original = service.ingest_parsed_orders

    async def fail_after_order_import(*args, **kwargs):
        await original(*args, **kwargs)
        raise RuntimeError("synthetic late failure")

    monkeypatch.setattr(service, "ingest_parsed_orders", fail_after_order_import)
    h, o = pair()
    with pytest.raises(RuntimeError, match="synthetic"):
        await service.import_pair(session, h, o, as_of=DAY)
    assert await row_counts(session) == [0, 0, 0, 0]


async def test_cancelled_pair_import_rolls_back_all_sections(db, monkeypatch):
    import asyncio

    session, _ = db
    original = service.ingest_parsed_orders

    async def cancel_after_orders(*args, **kwargs):
        await original(*args, **kwargs)
        raise asyncio.CancelledError()

    monkeypatch.setattr(service, "ingest_parsed_orders", cancel_after_orders)
    h, o = pair()
    with pytest.raises(asyncio.CancelledError):
        await service.import_pair(session, h, o, as_of=DAY)
    assert await row_counts(session) == [0, 0, 0, 0]


async def test_new_day_snapshot_is_not_suppressed_by_unchanged_export_bytes(db):
    session, _ = db
    h, o = pair()
    await service.import_pair(session, h, o, as_of=DAY)
    result = await service.import_pair(session, h, o, as_of=DAY + dt.timedelta(days=1))
    assert result["snapshot"] == "imported" and result["orders"] == "unchanged"
    assert await row_counts(session) == [2, 4, 1, 1]


async def test_missing_order_alias_cannot_match_another_account(db):
    session, _ = db
    from sqlalchemy import delete

    await session.execute(delete(AccountAlias))
    await session.commit()
    h, o = pair()
    with pytest.raises(service.BarclaysPairError, match="alias"):
        await service.import_pair(session, h, o, as_of=DAY)
    assert await row_counts(session) == [0, 0, 0, 0]


async def test_malformed_snapshot_cannot_close_preexisting_positions(db):
    session, _ = db
    h, o = pair()
    await service.import_pair(session, h, o, as_of=DAY)
    with pytest.raises(service.BarclaysPairError):
        await service.import_pair(
            session, change_cell(h, 1, None), o, as_of=DAY + dt.timedelta(days=1)
        )
    assert await row_counts(session) == [1, 2, 1, 1]
    from app.models import Instrument

    assert (
        await session.scalar(
            select(func.count()).select_from(Instrument).where(Instrument.closed_at.is_not(None))
        )
        == 0
    )


async def test_runner_imports_fetched_pair_before_acknowledging_success(db, tmp_path):
    from app.services.sync_runner import run_sync_all

    session, _ = db
    h, o = pair()
    acknowledged = []

    async def fetch(inbox):
        return service.FetchedPair(
            h, o, dt.datetime(2026, 9, 23, 18, 0, tzinfo=dt.UTC), lambda: acknowledged.append(True)
        )

    report = await run_sync_all(
        session, fetchers=[("Barclays", fetch)], inbox=tmp_path, include_trading212=False
    )
    assert report.steps[0].status == "ok"
    assert acknowledged == [True]
    assert await row_counts(session) == [1, 2, 1, 1]
    assert b"PK" not in (tmp_path / "last-sync.json").read_bytes()


async def test_failed_pair_import_never_acknowledges_login_attempt(db, tmp_path):
    from app.services.sync_runner import run_sync_all

    session, _ = db
    h, o = pair()
    acknowledged = []

    async def fetch(inbox):
        return service.FetchedPair(
            change_cell(h, 1, None),
            o,
            dt.datetime(2026, 9, 23, 18, 0, tzinfo=dt.UTC),
            lambda: acknowledged.append(True),
        )

    report = await run_sync_all(
        session, fetchers=[("Barclays", fetch)], inbox=tmp_path, include_trading212=False
    )
    assert report.steps[0].status == "failed"
    assert report.steps[0].detail == "BarclaysPairError"
    assert acknowledged == []
    assert await row_counts(session) == [0, 0, 0, 0]
